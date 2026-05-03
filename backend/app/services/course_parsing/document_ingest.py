from __future__ import annotations

from io import BytesIO
from pathlib import Path

from pydantic import BaseModel, Field, ValidationError

from app.schemas.course import DocumentType
from app.services.ai_gateway.service import AiGatewayError, generate_json


class ExtractedDocument(BaseModel):
    filename: str
    document_type: DocumentType = Field(alias="documentType")
    text: str


class DocumentTypePayload(BaseModel):
    document_type: DocumentType = Field(alias="documentType")


DOCUMENT_TYPE_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["documentType"],
    "properties": {
        "documentType": {
            "type": "string",
            "enum": [item.value for item in DocumentType],
        }
    },
}


async def extract_uploaded_document(filename: str, content: bytes) -> ExtractedDocument:
    text = extract_text(filename, content)
    document_type = classify_document_type(filename, text)
    return ExtractedDocument(filename=filename, documentType=document_type, text=text)


def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text(content)
    if suffix in {".pptx", ".ppt"}:
        return extract_presentation_text(content)
    if suffix == ".docx":
        return extract_docx_text(content)
    if suffix in {".md", ".markdown", ".txt"}:
        return decode_text(content)
    return decode_text(content)


def extract_pdf_text(content: bytes) -> str:
    import fitz

    with fitz.open(stream=content, filetype="pdf") as document:
        return "\n\n".join(page.get_text("text").strip() for page in document if page.get_text("text").strip())


def extract_presentation_text(content: bytes) -> str:
    from pptx import Presentation

    try:
        presentation = Presentation(BytesIO(content))
    except Exception:
        return decode_text(content)

    slides: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text.strip())
        if parts:
            slides.append(f"Slide {index}\n" + "\n".join(parts))
    return "\n\n".join(slides)


def extract_docx_text(content: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n\n".join(paragraphs)


def decode_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "gb18030", "latin-1"):
        try:
            return content.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore").strip()


def classify_document_type(filename: str, text: str) -> DocumentType:
    try:
        return classify_document_type_with_ai(filename, text)
    except (AiGatewayError, ValidationError, ValueError):
        return classify_document_type_fallback(filename, text)


def classify_document_type_with_ai(filename: str, text: str) -> DocumentType:
    prompt = f"""
Classify this uploaded course document into exactly one EduMind DocumentType.

Allowed values:
- syllabus: course outline, assessment policy, weekly schedule, learning outcomes
- slides: lecture slides, slide decks, presentation pages
- notes: lecture notes, reading notes, summaries
- past_paper: past exam paper or practice exam questions
- marking_scheme: answer key, rubric, grading guide, marking scheme
- tutorial: tutorial worksheet, lab sheet, discussion guide
- homework: assignment, problem set, take-home task
- quiz: quiz, short test, checkpoint questions
- other: anything else

Filename: {filename}

Content sample:
{text[:4000]}
"""
    raw = generate_json(
        prompt,
        DOCUMENT_TYPE_SCHEMA,
        system_prompt="You classify course documents. Return JSON only.",
        temperature=0,
    )
    payload = DocumentTypePayload.model_validate(raw)
    return payload.document_type


def classify_document_type_fallback(filename: str, text: str) -> DocumentType:
    haystack = f"{filename}\n{text[:2000]}".lower()
    if any(keyword in haystack for keyword in ("syllabus", "course outline", "learning outcomes")):
        return DocumentType.SYLLABUS
    if any(keyword in haystack for keyword in ("slide", "ppt", "lecture deck", "presentation")):
        return DocumentType.SLIDES
    if any(keyword in haystack for keyword in ("past paper", "past exam", "final exam", "exam paper")):
        return DocumentType.PAST_PAPER
    if any(keyword in haystack for keyword in ("marking scheme", "answer key", "rubric", "solution")):
        return DocumentType.MARKING_SCHEME
    if any(keyword in haystack for keyword in ("tutorial", "worksheet", "lab")):
        return DocumentType.TUTORIAL
    if any(keyword in haystack for keyword in ("homework", "assignment", "problem set")):
        return DocumentType.HOMEWORK
    if any(keyword in haystack for keyword in ("quiz", "checkpoint", "short test")):
        return DocumentType.QUIZ
    if Path(filename).suffix.lower() in {".ppt", ".pptx"}:
        return DocumentType.SLIDES
    if Path(filename).suffix.lower() in {".md", ".markdown", ".txt", ".docx", ".pdf"}:
        return DocumentType.NOTES
    return DocumentType.OTHER
